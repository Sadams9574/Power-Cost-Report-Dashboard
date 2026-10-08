"""Extract the Henderson baseline (Sheet1 hours, Weights curves, holiday factors)
into data/baseline.json for curves.html.

Usage: python3 scripts/extract_baseline.py [Henderson_BL_Creation.xlsx] [data/baseline.json]

Start/Finish come from Sheet1 columns G/H (hard-coded dates), hours from column I.
The weekly spread is re-computed with the workbook's own Weights formula and
checked against the cached Sheet1 values so the page and the workbook agree.
"""
import datetime as dt
import json
import math
import sys

import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else "Henderson_BL_Creation.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "data/baseline.json"
FIRST_WEEK_COL = 10      # Sheet1/Weights column J
FIRST_ROW, LAST_ROW = 4, 138
FACTOR_ROW = 140         # Weights row holding the holiday factor per week
HOURS_PER_FTE = 60       # Sheet1 FTE block divides weekly hours by 60


def iso(v):
    return v.date().isoformat() if isinstance(v, dt.datetime) else None


def num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def workbook_weights(weeks, factors, start, finish, curve):
    """Port of the Weights!J4 formula (Standard 3-week ramp, or Bell)."""
    s, f = dt.date.fromisoformat(start), dt.date.fromisoformat(finish) + dt.timedelta(days=6)
    active = [i for i, w in enumerate(weeks) if s <= w <= f and factors[i] > 0]
    n = len(active)
    out = [0.0] * len(weeks)
    for k, i in enumerate(active, 1):
        if curve == "Bell":
            w = math.exp(-0.5 * ((k - (n + 1) / 2) / max(n / 6, 0.5)) ** 2)
        else:
            w = 1 if n <= 6 else min(1, k / 3, (n - k + 1) / 3)
        out[i] = factors[i] * w
    return out


wb = openpyxl.load_workbook(SRC, data_only=True)
sh, wt = wb["Sheet1"], wb["Weights"]

week_cols = [c for c in range(FIRST_WEEK_COL, sh.max_column + 1) if isinstance(sh.cell(1, c).value, dt.datetime)]
weeks = [sh.cell(1, c).value.date() for c in week_cols]
factors = [num(wt.cell(FACTOR_ROW, c).value) if wt.cell(FACTOR_ROW, c).value is not None else 1.0 for c in week_cols]

codes, mismatches, missing_dates = [], [], []
for r in range(FIRST_ROW, LAST_ROW + 1):
    code = sh.cell(r, 2).value
    hours = num(sh.cell(r, 9).value)
    if not code or hours <= 0:
        continue
    start, finish = iso(sh.cell(r, 7).value), iso(sh.cell(r, 8).value)
    curve = (wt.cell(r, 3).value or "Standard").strip()
    cached = [num(sh.cell(r, c).value) for c in week_cols]
    entry = {
        "row": r,
        "category": sh.cell(r, 1).value,
        "code": str(code),
        "section": sh.cell(r, 3).value or "Other",
        "activity": str(sh.cell(r, 4).value or code).strip(),
        "start": start,
        "finish": finish,
        "hours": hours,
        "curve": curve,
    }
    if not (start and finish):
        missing_dates.append(entry["code"])
    else:
        w = workbook_weights(weeks, factors, start, finish, curve)
        tot = sum(w)
        calc = [hours * x / tot if tot else 0 for x in w]
        if any(abs(a - b) > 0.01 for a, b in zip(calc, cached)):
            mismatches.append(entry["code"])
    codes.append(entry)

data = {
    "source": SRC,
    "project": "Henderson",
    "hoursPerFTE": HOURS_PER_FTE,
    "weeks": [w.isoformat() for w in weeks],
    "holidayFactor": factors,
    "codes": codes,
}
with open(OUT, "w") as f:
    json.dump(data, f, indent=1)

total = sum(c["hours"] for c in codes)
print(f"Wrote {len(codes)} cost codes, {total:,.0f} hrs, {len(weeks)} weeks to {OUT}")
print(f"Bell curves: {[c['code'] for c in codes if c['curve'] == 'Bell']}")
print(f"Missing start/finish: {missing_dates or 'none'}")
print(f"Weekly spread vs workbook: {'all match' if not mismatches else 'MISMATCH ' + ', '.join(mismatches)}")
